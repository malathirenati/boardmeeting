import re
"""The board-meeting sheet format, in both directions.

One workbook per board meeting. Tabs: Guide, Meeting, then one tab per circle.
Each circle tab has one row per field:

  Panel | Row type | Label or text | Value | Value 2 | Value 3 | Value 4 | Value 5 | Value 6 | Comment

Rows with the same Panel code make one panel; panels appear in the order their code first appears.
"""
import datetime as dt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

CIRCLES = [("overview", "Overview", "OV"), ("ps", "Policy School", "PS"), ("rs", "Research", "RS"),
           ("md", "Media", "MD"), ("nw", "Network", "NW"), ("fn", "Finance", "FN")]
TYPES = {"text": "Figures and text", "line": "Line chart", "trend": "Trend chart", "bar": "Bar chart",
         "stacked": "Stacked bar chart", "carousel": "Pictures"}
TYPES_BACK = {v.lower(): k for k, v in TYPES.items()}
WIDTHS = {"full": "Full", "wide": "Wide", "narrow": "Narrow", "half": "Half", "third": "Third"}
ROW_TYPES = ["Introduction", "Panel", "Label above", "Width", "Unit", "Bars", "Source", "Note", "Needs checking", "Show",
             "Figure", "Point", "Paragraph", "Chart header", "Chart row", "Picture"]
HEAD = ["Panel", "Row type", "Label or text", "Value", "Value 2", "Value 3", "Value 4", "Value 5", "Value 6", "Comment (not shown)"]
NSER = 6

# ---------------------------------------------------------------- content -> rows
def panel_rows(code, p):
    R = []
    kind = p["chartType"] if p["type"] == "chart" else p["type"]
    R.append([code, "Panel", p["title"], TYPES[kind]])
    if p.get("eyebrow"): R.append([code, "Label above", p["eyebrow"]])
    R.append([code, "Width", WIDTHS.get(p.get("width", "full"), "Full")])
    if p["type"] == "chart":
        if p.get("unit"): R.append([code, "Unit", p["unit"]])
        if kind in ("bar", "stacked"): R.append([code, "Bars", "Horizontal" if p.get("horizontal") else "Vertical"])
    if p.get("source"): R.append([code, "Source", p["source"]])
    if p.get("note"): R.append([code, "Note", p["note"]])
    if p.get("flagged"): R.append([code, "Needs checking", "Yes"])
    if p["type"] == "text":
        for k in p.get("kpis", []): R.append([code, "Figure", k["label"], k["value"], k.get("note", "")])
        for block in [b for b in (p.get("body") or "").split("\n\n") if b.strip()]:
            lines = [l for l in block.split("\n") if l.strip()]
            if all(l.lstrip().startswith(("- ", "• ")) for l in lines):
                for l in lines: R.append([code, "Point", l.lstrip()[2:].strip()])
            else:
                R.append([code, "Paragraph", " ".join(lines)])
    elif p["type"] == "chart":
        t = p["table"]
        R.append([code, "Chart header"] + t[0])
        for r in t[1:]: R.append([code, "Chart row"] + [_num(c) for c in r[:1]] + [_num(c) for c in r[1:]])
    else:
        for s in p.get("slides", []): R.append([code, "Picture", s["caption"], s["img"]])
    return R

def _num(v):
    try:
        f = float(str(v).replace(",", ""))
        return int(f) if f.is_integer() and "." not in str(v) else f
    except ValueError:
        return v

def tab_rows(prefix, tab):
    R = [["TAB", "Introduction", tab.get("intro", "")]]
    for i, p in enumerate(tab["panels"], 1):
        R += panel_rows(f"{prefix}.{i:02d}", p)
    return R

# ---------------------------------------------------------------- write workbook
F = "Arial"
WINE, SOFT, PALE, DEEP, INK70, TEAL = "620D3C", "F5E6EC", "FCF0D9", "F1EEEA", "5C5755", "2F6B6B"
thin = Side(style="thin", color="D9D4CF")
SETTINGS = ["Label above", "Width", "Unit", "Bars", "Source", "Note", "Needs checking", "Show"]
HELP = {
 "Introduction": "C: one sentence shown under the tab heading on the dashboard.",
 "Panel": "Starts a panel. C: the panel title. D: what to show: Figures and text, Line chart, Trend chart, Bar chart, Stacked bar chart or Pictures.",
 "Label above": "C: small label above the title, e.g. the period or unit. Optional.",
 "Width": "C: Full, Wide, Narrow, Half or Third. A Wide and a Narrow panel sit side by side.",
 "Unit": "C: unit of the chart's numbers, e.g. ₹ lakh, students, %. Charts only.",
 "Bars": "C: Horizontal or Vertical. Bar and stacked bar charts only.",
 "Source": "C: where the figures come from. Shown under the panel.",
 "Note": "C: a note or caveat shown under the panel.",
 "Needs checking": "C: Yes flags the panel as not yet verified. No or blank otherwise.",
 "Show": "C: No hides the whole panel without deleting it. Yes or blank shows it.",
 "Figure": "C: what the figure is. D: the figure exactly as it should read (320, ₹37.72L, 83%). E: one line of context. Optional.",
 "Point": "C: one numbered point, in one sentence.",
 "Paragraph": "C: a paragraph of text.",
 "Chart header": "C: name of the categories (e.g. Cohort). D to I: one series name per column (e.g. DFA, APP, TP). One per chart, above its rows.",
 "Chart row": "C: the category (e.g. Jan 26). D to I: plain numbers, one per series. No ₹, %, commas or text.",
 "Picture": "C: caption. D: insert the picture straight into this cell (Insert → Image → Image in cell); Sync moves it to the Pictures folder. Or type the file name of a picture already in the Pictures folder.",
}
COL_NOTES = {
 1: "PANEL CODE\nRows with the same code make one panel on the dashboard, e.g. PS.03.\nNew panel: use the next free number (PS.09, PS.10 …).\nThe first row of the tab uses TAB.",
 2: "ROW TYPE\nPick from the drop-down. It decides what the row means. Column K then explains what to type.",
 3: "LABEL OR TEXT\nThe title, label, sentence or category, depending on the row type.",
 4: "VALUE\nThe figure, the panel type, the first series, or the picture link, depending on the row type.",
 5: "VALUE 2 TO VALUE 6\nUsed by charts (one column per series) and by Figure rows (context line in Value 2).",
 10: "YOUR COMMENT\nNotes for yourself or other leads. Never shown on the dashboard.",
 11: "WHAT TO ENTER\nFills in automatically from the row type. Do not type here.",
}
FIRST_NOTES = {
 "Panel": (4, "The panel type. Change it to switch, e.g. from Bar chart to Line chart. The rows below this one, up to the next Panel code, belong to this panel."),
 "Figure": (4, "Type the figure exactly as it should read on the dashboard: 320, ₹37.72L, 83%, 9.41."),
 "Point": (3, "Each Point row becomes one numbered point. Add a point: insert a row, copy the panel code, choose Point."),
 "Chart header": (4, "Series names go across D to I, one per column. The chart draws one line or bar colour per series."),
 "Chart row": (4, "Numbers only. To add a period, insert a Chart row below the last one and type the category in C."),
 "Picture": (4, "Click this cell, then Insert → Image → Image in cell, and choose the picture. At the next Sync it is saved to the Pictures folder and its file name appears here. Pictures appear as a slideshow in row order."),
 "Show": (3, "Set to No to hide this panel for this meeting without deleting the rows."),
}
VALUE_LISTS = {"Width": '"Full,Wide,Narrow,Half,Third"', "Bars": '"Horizontal,Vertical"', "Needs checking": '"Yes,No"', "Show": '"Yes,No"'}

def _fill(c): return PatternFill("solid", fgColor=c)

def _head(ws, cols, widths, row=1, freeze=True):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font = Font(name=F, bold=True, color="FFFFFF", size=10); x.fill = _fill(WINE)
        x.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[x.column_letter].width = w
    ws.row_dimensions[row].height = 30
    if freeze: ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate

def _note(cell, text, w=300, h=150):
    c = Comment(text, "Dashboard guide"); c.width, c.height = w, h; cell.comment = c

def _circle_tab(wb, name, prefix, rows, meeting_id):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = WINE if name != "Overview" else "F1A222"
    H = 4                                     # header row; data from row 5
    widths = [10, 16, 56, 18, 14, 14, 14, 14, 14, 26, 60]
    head = HEAD + ["What to enter (automatic)"]
    # banner
    ws.merge_cells("A1:K1"); ws["A1"] = f"{name}  ·  Takshashila Board Meeting {meeting_id}"
    ws["A1"].font = Font(name=F, size=14, bold=True, color="FFFFFF"); ws["A1"].fill = _fill(WINE); ws.row_dimensions[1].height = 28
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.merge_cells("A2:K2")
    ws["A2"] = ("Update the yellow cells. Each dark band starts a panel on the dashboard; the rows under it, with the same code in column A, fill that panel. "
                "Column K explains every row. To add a field: right-click a row in the panel → Insert 1 row below → copy the panel code into A → pick a row type in B. "
                "To remove a field, delete its row. Full instructions with examples: Guide tab.")
    ws["A2"].font = Font(name=F, size=10, color="171413"); ws["A2"].fill = _fill(PALE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center", indent=1); ws.row_dimensions[2].height = 48
    legend = [("Colour key:", None, "171413", True), ("Panel (starts a panel)", WINE, "FFFFFF", True), ("Setting (rarely changes)", DEEP, INK70, False),
              ("Content row", "FFFFFF", "171413", False), ("Cells to update", PALE, "171413", True), ("Chart series names", SOFT, WINE, True)]
    for j, (t, f, fc, b) in enumerate(legend, 1):
        c = ws.cell(row=3, column=[1, 3, 4, 6, 8, 10][j - 1] if j > 1 else 1, value=t)
        c.font = Font(name=F, size=9, bold=b, color=fc)
        if f: c.fill = _fill(f)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin) if f else Border()
        c.alignment = Alignment(vertical="center", indent=1)
    for a, b in (("D3", "E3"), ("F3", "G3"), ("H3", "I3"), ("J3", "K3")): ws.merge_cells(f"{a}:{b}")
    ws.merge_cells("A3:B3"); ws.row_dimensions[3].height = 22
    _head(ws, head, widths, row=H, freeze=False)
    ws.freeze_panes = "C5"
    for col, text in COL_NOTES.items(): _note(ws.cell(row=H, column=col), text)

    first = H + 1; seen = set(); dvs = {}
    for i, row in enumerate(rows, first):
        for j, v in enumerate(row, 1):
            x = ws.cell(row=i, column=j, value=v); x.font = Font(name=F, size=10); x.border = Border(bottom=thin)
            x.alignment = Alignment(wrap_text=(j == 3), vertical="top")
        rt = row[1]
        if rt in FIRST_NOTES and rt not in seen:
            col, text = FIRST_NOTES[rt]; _note(ws.cell(row=i, column=col), text, 280, 110); seen.add(rt)
        if rt == "Panel": dvs.setdefault("Panel", []).append(f"D{i}")
        if rt in VALUE_LISTS: dvs.setdefault(rt, []).append(f"C{i}")
    for rt, cells in dvs.items():
        f1 = '"' + ",".join(TYPES.values()) + '"' if rt == "Panel" else VALUE_LISTS[rt]
        dv = DataValidation(type="list", formula1=f1, allow_blank=rt != "Panel")
        if rt == "Panel": dv.prompt = "Choose what this panel shows."; dv.showInputMessage = True
        ws.add_data_validation(dv)
        for c in cells: dv.add(c)
    last = first + len(rows) + 60                     # room for new rows
    # one formula fills column K for every row, including rows leads add later (Google Sheets ARRAYFORMULA)
    k = ws.cell(row=first, column=11, value=f'=ARRAYFORMULA(IF(B{first}:B="","",IFERROR(VLOOKUP(B{first}:B,Lists!$A$1:$B$40,2,FALSE),"Unknown row type: pick one from the drop-down in column B.")))')
    for i in range(first, last + 1):
        x = ws.cell(row=i, column=11); x.font = Font(name=F, size=9, italic=True, color=INK70); x.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(ROW_TYPES)}", allow_blank=True, showErrorMessage=True,
                        error="Pick a row type from the drop-down. The Guide tab explains each one.", errorTitle="Row type")
    dv.prompt = "Choose what this row holds. Column K then shows what to type."; dv.promptTitle = "Row type"; dv.showInputMessage = True
    ws.add_data_validation(dv); dv.add(f"B{first}:B{last}")
    rng = lambda a, b: f"{a}{first}:{b}{last}"
    r = f"$B{first}"
    cf = ws.conditional_formatting
    cf.add(rng("A", "K"), FormulaRule(formula=[f'{r}="Panel"'], fill=_fill(WINE), font=Font(bold=True, color="FFFFFF"), stopIfTrue=True))
    cf.add(rng("A", "J"), FormulaRule(formula=[f'OR({",".join(f"{r}=" + chr(34) + s + chr(34) for s in SETTINGS)})'], fill=_fill(DEEP), font=Font(color=INK70)))
    cf.add(rng("A", "J"), FormulaRule(formula=[f'{r}="Chart header"'], fill=_fill(SOFT), font=Font(bold=True, color=WINE)))
    cf.add(rng("C", "C"), FormulaRule(formula=[f'OR({r}="Introduction",{r}="Point",{r}="Paragraph",{r}="Picture")'], fill=_fill(PALE)))
    cf.add(rng("D", "E"), FormulaRule(formula=[f'{r}="Figure"'], fill=_fill(PALE), font=Font(bold=True)))
    cf.add(rng("D", "I"), FormulaRule(formula=[f'{r}="Chart row"'], fill=_fill(PALE)))
    cf.add(rng("D", "D"), FormulaRule(formula=[f'{r}="Picture"'], fill=_fill(PALE)))
    cf.add(rng("B", "B"), FormulaRule(formula=[f'AND({r}<>"",COUNTIF(Lists!$A$1:$A$40,{r})=0)'], fill=_fill("F8E6E7"), font=Font(bold=True, color="A3282D")))
    ws.auto_filter.ref = f"A{H}:K{first + len(rows) - 1}"
    return ws

def _guide(wb, meeting_id):
    g = wb.active; g.title = "Guide"; g.sheet_properties.tabColor = "F1A222"
    for col, w in zip("ABCDEFG", [16, 16, 44, 22, 14, 14, 50]): g.column_dimensions[col].width = w
    r = [1]
    def line(text, size=10, bold=False, color="171413", fill=None, height=None, merge=True):
        i = r[0]
        if merge: g.merge_cells(f"A{i}:G{i}")
        c = g.cell(row=i, column=1, value=text); c.font = Font(name=F, size=size, bold=bold, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        if fill: c.fill = _fill(fill)
        if height: g.row_dimensions[i].height = height
        r[0] += 1
    def h(text): r[0] += 1; line(text, 13, True, WINE, height=24)
    def example(rows, caption=None):
        if caption: line(caption, 9, False, INK70)
        for row in rows:
            i = r[0]
            for j, v in enumerate(row, 1):
                c = g.cell(row=i, column=j, value=v); c.font = Font(name=F, size=9); c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            rt = row[1] if len(row) > 1 else ""
            if rt == "Row type" or row[0] == "Row type":
                for j in range(1, len(row) + 1): g.cell(row=i, column=j).fill = _fill(WINE); g.cell(row=i, column=j).font = Font(name=F, size=9, bold=True, color="FFFFFF")
            elif rt == "Panel":
                for j in range(1, min(len(row), 6) + 1): g.cell(row=i, column=j).fill = _fill(WINE); g.cell(row=i, column=j).font = Font(name=F, size=9, bold=True, color="FFFFFF")
            elif rt in SETTINGS:
                for j in range(1, min(len(row), 6) + 1): g.cell(row=i, column=j).fill = _fill(DEEP); g.cell(row=i, column=j).font = Font(name=F, size=9, color=INK70)
            elif rt == "Chart header":
                for j in range(1, min(len(row), 6) + 1): g.cell(row=i, column=j).fill = _fill(SOFT); g.cell(row=i, column=j).font = Font(name=F, size=9, bold=True, color=WINE)
            elif rt == "Figure":
                for j in (4, 5): g.cell(row=i, column=j).fill = _fill(PALE)
            elif rt == "Chart row":
                for j in range(4, min(len(row), 6) + 1): g.cell(row=i, column=j).fill = _fill(PALE)
            elif rt in ("Point", "Paragraph", "Introduction"):
                g.cell(row=i, column=3).fill = _fill(PALE)
            elif rt == "Picture":
                for j in (3, 4): g.cell(row=i, column=j).fill = _fill(PALE)
            if len(row) > 6 and row[6] and rt != "Row type" and row[0] != "Row type": g.cell(row=i, column=7).font = Font(name=F, size=9, bold=True, color="2F6B4A")
            r[0] += 1
    HDR = ["Panel", "Row type", "Label or text", "Value", "Value 2", "Value 3", ""]

    g.merge_cells("A1:G1"); g["A1"] = "How to fill in your circle's tab"; g["A1"].font = Font(name=F, size=16, bold=True, color="FFFFFF"); g["A1"].fill = _fill(WINE)
    g["A1"].alignment = Alignment(vertical="center", indent=1); g.row_dimensions[1].height = 34; r[0] = 2
    line(f"This sheet is everything the dashboard shows for board meeting {meeting_id}. Each circle lead normally updates their own circle's tab, and anyone listed can edit any tab. "
         "You do not need to change anything else: the dashboard reads the rows and lays out the panels itself.", 10, height=44)

    h("1. Start here")
    for t in ["① Open your circle's tab at the bottom of the screen (Policy School, Research, Media, Network, Finance; Overview is shared).",
              "② Go down the tab panel by panel. Each dark band is one panel on the dashboard; the rows under it fill that panel.",
              "③ Change the yellow cells: figures, points, chart numbers and picture links. Column K always tells you what a row expects.",
              "④ Add, remove or hide fields as described below. Hover over any cell with a small black triangle in its corner for a tip.",
              "⑤ When your tab is done, press Sync on the dashboard (or Board dashboard → Sync the dashboard now in the index sheet). The dashboard updates in two to three minutes."]:
        line(t, height=30)

    h("2. Colour key")
    example([["PS.01", "Panel", "GCPP programme, May 2026 cohort", "Figures and text", "", "", "Dark band: starts a panel. Title in C, panel type in D."],
             ["PS.01", "Width", "Full", "", "", "", "Grey: a setting. Rarely changes."],
             ["PS.01", "Figure", "Total enrolled", "320", "APP 43 · DFA 244 · TP 36", "", "Yellow: cells you update each meeting."],
             ["PS.02", "Chart header", "Cohort", "DFA", "APP", "TP", "Pink: the chart's series names."],
             ["PS.02", "Chart row", "May 26", "244", "43", "36", "Yellow numbers: one per series."]])

    h("3. Change a figure or a sentence")
    line("Type over the yellow cell. Write figures exactly as they should appear: ₹37.72L, 83%, 1,250. Keep chart numbers plain: 37.72, 83, 1250.", height=30)

    h("4. Add a new field")
    line("Example: add 'Referral enrolments' to the GCPP panel (PS.01).", 10, True)
    line("Right-click the last row of the panel → Insert 1 row below. In the new row type PS.01 in A, choose Figure in B, then fill C, D and E.", height=30)
    example([HDR[:-1] + ["What it does"],
             ["PS.01", "Figure", "Average revenue per student", "₹10.8k", "Target ₹15k", "", ""],
             ["PS.01", "Figure", "Referral enrolments", "14", "Up from 12 in Feb 2026", "", "← new row: appears as a fifth figure"]])
    line("The same works for every kind of field: a Point row adds a numbered point; a Chart row adds a period to a chart; a Picture row adds a slide.", height=30)
    line("Add a picture: insert a row in the Pictures panel, choose Picture in B, type the caption in C, click D and choose Insert → Image → Image in cell. Sync saves it to the Pictures folder and puts its file name in D.", height=30)

    h("5. Add a new panel")
    line("Example: a bar chart of marketing spend. Use the next free code in your tab (here PS.09). Add the rows at the end of the tab, or where the panel should appear.", height=30)
    example([HDR[:-1] + ["What it does"],
             ["PS.09", "Panel", "Marketing spend by channel", "Bar chart", "", "", "Starts the panel and picks the chart type"],
             ["PS.09", "Label above", "₹ lakh · Apr–Jun 2026", "", "", "", "Optional"],
             ["PS.09", "Width", "Narrow", "", "", "", "Optional; Full if left out"],
             ["PS.09", "Unit", "₹ lakh", "", "", "", "Optional"],
             ["PS.09", "Chart header", "Channel", "Spend", "", "", "Series names across"],
             ["PS.09", "Chart row", "Search ads", "2.1", "", "", "One row per bar"],
             ["PS.09", "Chart row", "Social media", "1.4", "", "", ""],
             ["PS.09", "Source", "Policy School marketing tracker", "", "", "", "Shown under the chart"]])
    line("Panel types: Figures and text · Line chart · Trend chart (line plus a fitted trend) · Bar chart · Stacked bar chart · Pictures (a slideshow with captions).", height=30)

    h("6. Remove or hide")
    line("Remove a field: delete its row (right-click the row number → Delete row).  Hide a whole panel for this meeting but keep it for later: add a row  PS.04 · Show · No.", height=30)

    h("7. Move a panel")
    line("Panels appear in the order their code first appears in the tab. To move a panel, cut its rows and insert them where it should go. The codes do not need to be in number order.", height=30)

    h("8. Row types: full reference")
    example([["Row type", "", "What to type (column K shows the same text)", "", "", "", "Example"]] +
            [[t, "", HELP[t], "", "", "", ex] for t, ex in [
             ("Introduction", "Updates from the Policy School circle."), ("Panel", "GCPP programme · Figures and text"),
             ("Label above", "Q4 FY26 performance"), ("Width", "Wide"), ("Figure", "Total enrolled · 320 · APP 43 · DFA 244"),
             ("Point", "72% TP completion."), ("Paragraph", "A short paragraph."), ("Chart header", "Cohort · DFA · APP · TP"),
             ("Chart row", "May 26 · 244 · 43 · 36"), ("Unit", "students"), ("Bars", "Horizontal"), ("Picture", "Caption · picture inserted in the cell"),
             ("Source", "Takshashila board meeting deck, June 2026"), ("Note", "Streams add to 323."), ("Needs checking", "Yes"), ("Show", "No")]])
    for i in range(r[0] - 16, r[0]):
        g.merge_cells(f"C{i}:F{i}"); g.row_dimensions[i].height = 32
        g.cell(row=i, column=1).font = Font(name=F, size=9, bold=True)

    h("9. Rules that keep the dashboard working")
    for t in ["Do not rename the tabs or move the columns. Add rows, never columns.",
              "Every panel needs exactly one Panel row. Each code is used for one panel only within a tab.",
              "Chart rows take plain numbers only: no ₹, %, commas, 'K' or text. Leave a cell blank if there is no figure.",
              "Put pictures inside the cell (Insert → Image → Image in cell), not floating over the sheet: floating pictures are not picked up. Sync saves them to the Pictures folder.",
              "Column J is for your own comments; the dashboard never shows it."]:
        line("•  " + t, height=24)

    h("10. Each board meeting")
    line("The dashboard owner or any lead starts a new meeting from the index sheet: Board dashboard → Start a new board meeting. "
         "That copies this sheet, clears the figures and pictures, and keeps the panels, charts and text for you to update. Earlier meetings move to the dashboard's archive automatically.", height=44)

def write_meeting(content, path, guide_rows=None):
    m = content["meeting"]
    wb = Workbook()
    _guide(wb, m["id"])
    ws = wb.create_sheet("Meeting"); ws.sheet_properties.tabColor = "F1A222"
    _head(ws, ["Field", "Value", "What it means"], [28, 30, 80])
    rows = [("Meeting ID", m["id"], "Year and month of the meeting, yyyy-mm. Must be unique. Set automatically for a new meeting."),
            ("Meeting title", m.get("title", "Takshashila Board Meeting"), "Shown on the dashboard."),
            ("Meeting date", dt.date.fromisoformat(m["date"]), m.get("date_note", "") or "Date of the board meeting."),
            ("Reporting window from", dt.date.fromisoformat(m["window_start"]), "First day the figures cover. Any length of period works."),
            ("Reporting window to", dt.date.fromisoformat(m["window_end"]), "Last day the figures cover. Where a circle's figures cover different dates, say so in that panel's 'Label above'."),
            ("Deck label", m.get("deck_label", ""), "Optional: how the period is described, e.g. FY26 Q2.")]
    for i, (a, b, c) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=a).font = Font(name=F, size=10, bold=True)
        x = ws.cell(row=i, column=2, value=b); x.font = Font(name=F, size=10, color="0000FF"); x.fill = _fill(PALE)
        if isinstance(b, dt.date): x.number_format = "yyyy-mm-dd"
        y = ws.cell(row=i, column=3, value=c); y.font = Font(name=F, size=10, italic=True, color=INK70); y.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[i].height = 30
    lists = wb.create_sheet("Lists")
    for i, v in enumerate(ROW_TYPES, 1):
        lists.cell(row=i, column=1, value=v); lists.cell(row=i, column=2, value=HELP[v])
    lists.sheet_state = "hidden"
    for key, name, prefix in CIRCLES:
        _circle_tab(wb, name, prefix, tab_rows(prefix, content["tabs"][key]), m["id"])
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.save(path)

# ---------------------------------------------------------------- read workbook
def _group(v):
    """Sheets turns a typed 859,530 into the number 859530; put the separators back for display."""
    return f"{int(v):,}" if re.fullmatch(r"\d{4,}", v or "") else v

def _s(v):
    if v is None: return ""
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def read_meeting(path_or_file):
    wb = load_workbook(path_or_file, data_only=True)
    warn = []
    ms = wb["Meeting"]
    kv = {_s(r[0]).lower(): r[1] for r in ms.iter_rows(min_row=2, values_only=True) if r and r[0]}
    def date(k):
        v = kv.get(k)
        if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%Y-%m-%d")
        return _s(v)
    meeting = {"id": _s(kv.get("meeting id")), "title": _s(kv.get("meeting title")) or "Takshashila Board Meeting",
               "date": date("meeting date"), "window_start": date("reporting window from"), "window_end": date("reporting window to"),
               "deck_label": _s(kv.get("deck label"))}
    for k in ("id", "date", "window_start", "window_end"):
        if not meeting[k]: warn.append(f"Meeting tab: '{k}' is empty.")
    tabs = {}
    for key, name, prefix in CIRCLES:
        if name not in wb.sheetnames:
            tabs[key] = {"name": name, "intro": "", "panels": []}; warn.append(f"No '{name}' tab; shown empty."); continue
        tabs[key] = read_tab(wb[name], name, warn)
    return {"version": 3, "meeting": meeting, "tabs": tabs}, warn

def read_tab(ws, name, warn):
    intro, order, P = "", [], {}
    hdr = 1
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
        if len(row) > 1 and _s(row[1]).lower() == "row type": hdr = i; break
    for i, row in enumerate(ws.iter_rows(min_row=hdr + 1, values_only=True), hdr + 1):
        row = list(row) + [None] * 10
        code, rt, c = _s(row[0]), _s(row[1]).lower(), _s(row[2])
        vals = [_s(v) for v in row[3:3 + NSER]]
        if not rt: continue
        if rt == "introduction": intro = c; continue
        if not code: warn.append(f"{name} row {i}: no Panel code; row skipped."); continue
        if code not in P:
            order.append(code)
            P[code] = {"id": f"{name[:2].lower()}-{code}", "type": None, "title": "", "eyebrow": "", "width": "full", "source": "", "note": "",
                       "flagged": False, "show": True, "kpis": [], "body": [], "table": [], "slides": [], "unit": "", "horizontal": None, "_row": i}
        p = P[code]
        if rt == "panel":
            p["title"] = c; k = TYPES_BACK.get(vals[0].lower())
            if not k: warn.append(f"{name} row {i}: panel type '{vals[0]}' not recognised; shown as Figures and text."); k = "text"
            p["kind"] = k
        elif rt == "label above": p["eyebrow"] = c
        elif rt == "width": p["width"] = c.lower() if c.lower() in WIDTHS else "full"
        elif rt == "unit": p["unit"] = c
        elif rt == "bars": p["horizontal"] = c.lower().startswith("h")
        elif rt == "source": p["source"] = c
        elif rt == "note": p["note"] = c
        elif rt == "needs checking": p["flagged"] = c.lower() in ("yes", "y", "true")
        elif rt == "show": p["show"] = c.lower() not in ("no", "n", "false", "hide")
        elif rt == "figure": p["kpis"].append({"label": c, "value": _group(vals[0]), "note": vals[1]})
        elif rt == "point": p["body"].append(("pt", c))
        elif rt == "paragraph": p["body"].append(("para", c))
        elif rt == "chart header": p["table"].insert(0, [c] + [v for v in vals if v])
        elif rt == "chart row": p["table"].append([c] + vals)
        elif rt == "picture": p["slides"].append({"caption": c, "img": vals[0]})
        else: warn.append(f"{name} row {i}: row type '{row[1]}' not recognised; skipped.")
    panels = []
    for code in order:
        p = P[code]
        if not p["show"]: continue
        kind = p.get("kind")
        if not kind: warn.append(f"{name} {code}: no 'Panel' row; shown as Figures and text."); kind = "text"
        base = {k: p[k] for k in ("id", "title", "eyebrow", "width", "source", "note", "flagged")}
        base["code"] = code
        if kind == "text":
            body, prev = [], None
            for t, v in p["body"]:
                if t == "pt": body.append(("\n" if prev == "pt" else ("\n\n" if body else "")) + "- " + v)
                else: body.append(("\n\n" if body else "") + v)
                prev = t
            panels.append({**base, "type": "text", "kpis": p["kpis"], "body": "".join(body)})
        elif kind == "carousel":
            panels.append({**base, "type": "carousel", "slides": p["slides"]})
        else:
            t = p["table"]
            if not t: warn.append(f"{name} {code}: chart has no rows."); t = [["Category", "Value"]]
            w = len(t[0]); t = [(r + [""] * w)[:w] for r in t]
            hz = p["horizontal"] if p["horizontal"] is not None else kind in ("bar", "stacked")
            panels.append({**base, "type": "chart", "chartType": kind, "horizontal": hz, "unit": p["unit"], "table": t})
    return {"name": name, "intro": intro, "panels": panels}

# ---------------------------------------------------------------- index workbook
INDEX_HEAD = ["Meeting ID", "Meeting date", "Meeting sheet (link or file name)", "On dashboard", "Notes"]

def write_index(rows, path):
    wb = Workbook(); ws = wb.active; ws.title = "Meetings"; ws.sheet_properties.tabColor = WINE
    _head(ws, INDEX_HEAD, [14, 14, 70, 14, 60])
    notes = {1: "Year and month of the meeting, yyyy-mm.", 2: "Date of the board meeting.",
             3: "Link to the meeting's Google Sheet (Share → Copy link).",
             4: "Yes = shown on the dashboard. The latest meeting marked Yes is the current one; all earlier ones form the archive. Keep a new meeting at No until every circle has filled in its tab.",
             5: "Anything useful, e.g. who created it."}
    for c, t in notes.items(): _note(ws.cell(row=1, column=c), t)
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            x = ws.cell(row=i, column=j, value=v); x.font = Font(name=F, size=10)
            if isinstance(v, dt.date): x.number_format = "yyyy-mm-dd"
            if j == 4: x.fill = _fill(PALE)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False); ws.add_data_validation(dv); dv.add("D2:D200")
    n = ws.max_row + 2
    for k, t in enumerate(["Start a new meeting: menu Board dashboard → Start a new board meeting… (it copies the latest sheet, sets the dates, shares it with the circle leads and adds a row here).",
                           "Show it: set On dashboard to Yes, then press Sync on the dashboard or use Board dashboard → Sync the dashboard now. The dashboard also refreshes every night.",
                           "Everyone on the Circle leads tab can edit every tab of each meeting sheet."]):
        ws.cell(row=n + k, column=1, value=t).font = Font(name=F, size=10, italic=True, color=INK70)
    cl = wb.create_sheet("Circle leads"); cl.sheet_properties.tabColor = "F1A222"
    _head(cl, ["Circle", "Lead email addresses (comma-separated)", "Role"], [26, 70, 70])
    for i, (tab, what) in enumerate([("Meeting", "Dashboard owners: set up meetings."),
                                     ("Overview", "Headline figures and the big developments, usually the CEO's office."),
                                     ("Policy School", "Policy School lead(s)."), ("Research", "Research lead(s)."), ("Media", "Media lead(s)."),
                                     ("Network", "Operations and alumni lead(s)."), ("Finance", "Finance lead(s).")], 2):
        cl.cell(row=i, column=1, value=tab).font = Font(name=F, size=10, bold=True)
        x = cl.cell(row=i, column=2); x.fill = _fill(PALE); x.font = Font(name=F, size=10, color="0000FF")
        cl.cell(row=i, column=3, value=what).font = Font(name=F, size=10, italic=True, color=INK70)
    cl.cell(row=10, column=1, value="Everyone listed gets edit access to each new meeting sheet and can edit any tab; no tab is locked. Google Sheets keeps a version history (File → Version history) showing who changed what. "
            "After adding someone here, run Board dashboard → Share a meeting with circle leads to give them access to an existing meeting sheet.").font = Font(name=F, size=10, italic=True, color=INK70)
    _note(cl.cell(row=1, column=2), "Use the Google accounts people sign in with, e.g. name@takshashila.org.in. Separate several addresses with commas.")
    wb.save(path)

def read_index(path_or_file):
    wb = load_workbook(path_or_file, data_only=True); ws = wb["Meetings"]; out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or not r[2]: continue
        if not _s(r[0])[:4].isdigit(): continue
        out.append({"id": _s(r[0]), "date": _s(r[1]), "sheet": _s(r[2]), "on": _s(r[3]).lower() in ("yes", "y", "true")})
    return out
